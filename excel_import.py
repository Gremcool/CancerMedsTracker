import io
import re
import pandas as pd
from database import get_connection

# NOTE: save_medicines lives ONLY in database.py. A second copy used to live here and,
# because app.py star-imports both modules, it silently shadowed the safe version.

# ----------------------------------------------------------------------------
# FIELDS THE DIRECTORY UNDERSTANDS (same columns the Medicines grid has)
# ----------------------------------------------------------------------------
STATUS_OPTIONS = ["Open", "In Progress", "Waiting Supplier", "Escalated", "Completed"]

# header (as written in the template) -> (internal field, required?, help text)
TEMPLATE_BASE_COLUMNS = [
    ("ITEM DESCRIPTION", "medicine_name", True,  "Full item name exactly as you want it stored. Required."),
    ("BASE DRUG NAME",   "base_drug_name", False, "Optional grouping name. Leave blank to auto-detect from the item name."),
    ("STOCK ON HAND",    "stock_on_hand", False, "Optional number (SOH). Blank = 0 for new items."),
    ("IN TRANSIT",       "in_transit",    False, "Optional number (inbound / in transit)."),
    ("AMC",              "amc",           False, "Optional number (average monthly consumption)."),
    ("REQUEST QTY",      "request_qty",   False, "Optional number (quantity to request / order)."),
    ("STATUS",           "status",        False, "Optional: Open, In Progress, Waiting Supplier, Escalated or Completed."),
    ("OWNER",            "owner",         False, "Optional: person responsible."),
    ("NOTE",             "note",          False, "Optional: first log note for the item."),
]
NUMERIC_FIELDS = ["stock_on_hand", "in_transit", "amc", "request_qty"]
TRACKING_FIELDS = ["status", "owner", "note"]
RESERVED_FIELDS = {"medicine_name", "base_drug_name", "category", "import_status"} | set(NUMERIC_FIELDS) | set(TRACKING_FIELDS)

# (field, label shown in the mapping panel, required?)
IMPORT_FIELDS = [
    ("medicine_name", "Item description", True),
    ("base_drug_name", "Base drug / group name", False),
    ("stock_on_hand", "Stock on hand (SOH)", False),
    ("in_transit", "In transit / inbound", False),
    ("amc", "AMC", False),
    ("request_qty", "Request quantity", False),
    ("status", "Status", False),
    ("owner", "Owner", False),
    ("note", "Note / comment", False),
]

# Patterns are tried in order; the first sheet column that matches (and is not already used) wins.
_DETECT_PATTERNS = {
    "medicine_name": [r"^ITEM DESCRIPTION$", r"^MEDICINE NAME$", r"^ITEM NAME$", r"^ITEM$", r"^DESCRIPTION$", r"^NAME$"],
    "base_drug_name": [r"^BASE DRUG NAME$", r"^DRUG GROUP$", r"^GROUP( NAME)?$"],
    "stock_on_hand": [r"^(TOTAL )?(SOH|STOCK ON HAND)\b", r"^STOCK$"],
    "in_transit": [r"^(IN ?TRANSIT|INBOUND|TRANSIT|ON ORDER)\b"],
    "amc": [r"^AMC\b"],
    "request_qty": [r"^REQUEST(ED)? (QTY|QUANTIT)", r"^(QTY|QUANTITY) TO (ORDER|REQUEST)", r"^ORDER QTY"],
    "status": [r"^STATUS$"],
    "owner": [r"^(OWNER|RESPONSIBLE( OFFICER)?|ASSIGNED TO)$"],
    "note": [r"^(NOTE|NOTES|COMMENT|COMMENTS)$"],
}


def _norm(header):
    return re.sub(r"\s+", " ", str(header)).strip().upper()


_GROUP_TAIL_WORDS = {"SMALL", "MEDIUM", "LARGE", "XL", "XXL", "XS", "SIZE", "NO", "NO.", "NUMBER",
                     "WITH", "AND", "FOR", "OF", "IN", "TO", "X"}


def _clean_for_group(name):
    """Drops (...) notes and everything after the first comma, e.g. 'GLOVES, NITRILE (ASTM D6978), LARGE' -> 'GLOVES'."""
    text = str(name).strip()
    text = re.sub(r"\([^)]*\)?", " ", text)   # (...) including an unclosed "("
    text = text.split(",")[0]
    return re.sub(r"\s+", " ", text).strip()


def extract_base_drug_name(full_name):
    """
    Group name for an item: the product / drug name without strength, size or model code.
      'Bosutinib Tab. 400 mg'                     -> 'BOSUTINIB TAB.'
      'CHEMOTHERAPY GLOVES, NITRILE (...), LARGE'  -> 'CHEMOTHERAPY GLOVES'
      'CONNECTOR C35'                              -> 'CONNECTOR'
    """
    if pd.isna(full_name):
        return "UNKNOWN"
    raw = str(full_name).strip()
    text = _clean_for_group(raw) or raw
    match = re.search(r"\d", text)
    if match:
        base = text[:match.start()]
        if base and not base.endswith(" "):      # letters glued to the number are a model code (C35, U100, NO.15)
            base = re.sub(r"\S+$", "", base)
        base = re.sub(r"[\s,\-/]+$", "", base)
        if not base:                              # name starts with a number -> first word
            words = text.split()
            return words[0].strip(",;").upper() if words else "UNKNOWN"
    else:
        base = text
    tokens = base.split()
    while len(tokens) > 1 and tokens[-1].upper() in _GROUP_TAIL_WORDS:
        tokens.pop()
    return " ".join(tokens).strip(",;-/ ").upper() or "UNKNOWN"


def legacy_group_name(full_name):
    """The group name the app used before smart grouping existed (text before the first number).
    Used to tell automatically made groups from groups someone typed by hand."""
    text = str(full_name).strip()
    match = re.search(r'\d+', text)
    if match:
        base = re.sub(r'[\s,\-]+$', '', text[:match.start()].strip())
        if base:
            return base.upper()
    words = text.split()
    return words[0].strip(",;").upper() if words else "UNKNOWN"


def assign_groups(names, existing_groups=None):
    """
    Group name for every item in 'names'. Similar items end up under ONE group:
    a name that starts with another group's name joins it ('CHEMO PORT NEEDLE' -> 'CHEMO PORT'),
    and plural / singular variants are merged ('NEEDLES' -> 'NEEDLE').
    'existing_groups' are the groups already in the category, so new items join them.
    """
    bases = [extract_base_drug_name(n) for n in names]
    pool = {b for b in bases if b != "UNKNOWN"} | {str(g).strip().upper() for g in (existing_groups or []) if str(g).strip()}

    def merge(base):
        words = base.split()
        for i in range(1, len(words)):
            candidate = " ".join(words[:i])
            if candidate in pool:
                return candidate
        if base.endswith("S") and len(base) > 3 and base[:-1] in pool:
            return base[:-1]
        return base

    return [merge(b) for b in bases]


def load_excel_sheets(uploaded_file):
    xl = pd.ExcelFile(uploaded_file)
    return xl.sheet_names


def load_excel(uploaded_file, sheet_name=0):
    return pd.read_excel(uploaded_file, sheet_name=sheet_name)


# ----------------------------------------------------------------------------
# TEMPLATE HELPERS
# ----------------------------------------------------------------------------
def get_template_columns(custom_cols=None):
    """Ordered list of column headers a sheet for this category should use."""
    return [h for h, _, _, _ in TEMPLATE_BASE_COLUMNS] + list(custom_cols or [])


def get_template_guide(custom_cols=None):
    """Small table describing each column (order, name, required, notes) for on-screen display."""
    rows = []
    for i, (header, _, required, note) in enumerate(TEMPLATE_BASE_COLUMNS, start=1):
        rows.append({"Order": i, "Column header": header, "Required": "Yes" if required else "No", "Notes": note})
    offset = len(TEMPLATE_BASE_COLUMNS)
    for j, col in enumerate(custom_cols or [], start=1):
        rows.append({"Order": offset + j, "Column header": col, "Required": "No", "Notes": "Custom column for this category."})
    return pd.DataFrame(rows)


def build_template_excel(category, custom_cols=None):
    """Builds a ready-to-fill .xlsx template (header row only) and returns its bytes."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Template"

    notes = {h: n for h, _, _, n in TEMPLATE_BASE_COLUMNS}
    headers = get_template_columns(custom_cols)
    fill = PatternFill("solid", start_color="1F4E78")
    for idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=idx, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.comment = Comment(notes.get(header, f"Custom column for {category}."), "Template")
        ws.column_dimensions[get_column_letter(idx)].width = 48 if idx == 1 else 18
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_blank_entry_df(custom_cols=None, rows=5):
    """Empty text grid in the template column order ('type it in'). All text so empty cells show blank, not 'None'."""
    return pd.DataFrame({header: [""] * rows for header in get_template_columns(custom_cols)})


def parse_pasted_text(text):
    """
    Turns text pasted from Excel (tab separated) or a CSV into a DataFrame.
    * Tab separated  -> first line is the header row.
    * Comma header   -> used when the first line contains 'ITEM DESCRIPTION'.
    * Anything else  -> treated as a plain list: one item name per line.
    """
    text = (text or "").strip("\n\r ")
    if not text.strip():
        return pd.DataFrame()
    lines = [ln for ln in text.splitlines() if ln.strip()]
    if "\t" in text:
        return pd.read_csv(io.StringIO("\n".join(lines)), sep="\t", dtype=str, keep_default_na=False)
    if "ITEM DESCRIPTION" in lines[0].upper() and "," in lines[0]:
        return pd.read_csv(io.StringIO("\n".join(lines)), sep=",", dtype=str, keep_default_na=False)
    return pd.DataFrame({"ITEM DESCRIPTION": [ln.strip() for ln in lines]})


# ----------------------------------------------------------------------------
# COLUMN DETECTION + MAPPING
# ----------------------------------------------------------------------------
def detect_column_mapping(df):
    """
    Works out which sheet column feeds which field, e.g. 'TOTAL SOH 03,JUNE 2026' -> stock_on_hand,
    'QTY TO ORDER' -> request_qty. Returns {field: sheet column name (str) or None}.
    """
    mapping = {field: None for field, _, _ in IMPORT_FIELDS}
    if df is None or len(df.columns) == 0:
        return mapping

    columns = [str(c) for c in df.columns]
    normed = {c: _norm(c) for c in columns}
    used = set()

    for field, _, _ in IMPORT_FIELDS:
        found = None
        for pattern in _DETECT_PATTERNS[field]:
            found = next((c for c in columns if c not in used and re.search(pattern, normed[c])), None)
            if found:
                break
        # last resort for the item column (original behaviour): anything containing DESCRIPTION / ITEM
        if field == "medicine_name" and not found:
            found = next((c for c in columns if c not in used and ("DESCRIPTION" in normed[c] or "ITEM" in normed[c])), None)
        if found:
            mapping[field] = found
            used.add(found)
    return mapping


def clean_extra_name(col):
    """Header text used for a sheet column that is brought along as a custom column."""
    name = " ".join(str(col).split())
    if name.lower() in RESERVED_FIELDS:
        name = f"{name} (sheet)"
    return name


def _cell_text(value):
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return "" if text.lower() == "nan" else text


def _to_number(series):
    cleaned = series.astype(str).str.replace(",", "", regex=False).str.strip()
    return pd.to_numeric(cleaned, errors="coerce").round(0)


# ----------------------------------------------------------------------------
# PREPARATION
# ----------------------------------------------------------------------------
def prepare_medicines(df, custom_cols=None, mapping=None, extra_cols=None, existing_groups=None):
    """
    Maps spreadsheet columns into the directory layout.

    * Item names are kept EXACTLY as written in the sheet (only trimmed).
    * 'mapping' is {field: sheet column}; when omitted it is auto-detected from the headers.
    * Sheet columns matching this category's existing custom columns are picked up automatically.
    * 'extra_cols' are additional sheet columns brought along as custom columns.
    """
    custom_cols = custom_cols or []
    extra_cols = extra_cols or []
    if df is None or df.empty:
        return pd.DataFrame()

    col_by_str = {str(c): c for c in df.columns}
    mapping = mapping if mapping is not None else detect_column_mapping(df)

    def src(field):
        name = mapping.get(field)
        return col_by_str.get(str(name)) if name is not None else None

    desc_col = src("medicine_name")
    if desc_col is None:
        return pd.DataFrame()

    out = pd.DataFrame({"medicine_name": df[desc_col]})
    used_sources = {str(desc_col)}

    base_col = src("base_drug_name")
    if base_col is not None:
        out["_base_given"] = df[base_col]
        used_sources.add(str(base_col))

    for field in NUMERIC_FIELDS:
        col = src(field)
        if col is not None:
            out[field] = _to_number(df[col])
            used_sources.add(str(col))

    # Tracking columns (what the Medicines grid shows as Status / Owner / Last Note)
    status_col, owner_col, note_col = src("status"), src("owner"), src("note")
    if status_col is not None or owner_col is not None or note_col is not None:
        valid = {s.lower(): s for s in STATUS_OPTIONS}
        n = len(df)
        status_raw = df[status_col].map(_cell_text) if status_col is not None else pd.Series([""] * n, index=df.index)
        owner_vals = df[owner_col].map(_cell_text) if owner_col is not None else pd.Series([""] * n, index=df.index)
        note_vals = df[note_col].map(_cell_text) if note_col is not None else pd.Series([""] * n, index=df.index)
        status_ok = status_raw.map(lambda t: valid.get(t.lower(), ""))
        # a status the app does not know is kept as part of the note instead of being lost
        note_final = [
            (f"Status from sheet: {raw}. {note}".strip() if raw and not ok else note)
            for raw, ok, note in zip(status_raw, status_ok, note_vals)
        ]
        if status_col is not None:
            out["status"] = status_ok
            used_sources.add(str(status_col))
        if owner_col is not None:
            out["owner"] = owner_vals
            used_sources.add(str(owner_col))
        if note_col is not None or status_col is not None:
            out["note"] = pd.Series(note_final, index=df.index)
            if note_col is not None:
                used_sources.add(str(note_col))

    # Custom columns: existing ones for this category + any extra sheet columns the user picked
    custom_added = []
    norm_to_src = {}
    for c in df.columns:
        norm_to_src.setdefault(_norm(c), c)
    for name in custom_cols:
        col = norm_to_src.get(_norm(name))
        if col is not None and str(col) not in used_sources and name not in RESERVED_FIELDS:
            out[name] = df[col].map(_cell_text)
            used_sources.add(str(col))
            custom_added.append(name)
    for extra in extra_cols:
        col = col_by_str.get(str(extra))
        if col is None or str(col) in used_sources:
            continue
        name = clean_extra_name(col)
        if name in out.columns:
            continue
        out[name] = df[col].map(_cell_text)
        used_sources.add(str(col))

    out = out.dropna(subset=["medicine_name"])
    out["medicine_name"] = out["medicine_name"].astype(str).str.strip()
    out = out[out["medicine_name"] != '']
    out = out[~out["medicine_name"].str.contains("COMPATIBILITY REPORT|RUN ON|TOTAL", case=False, na=False)]
    out = out[out["medicine_name"].str.upper() != "NAN"]

    if out.empty:
        return pd.DataFrame()

    extracted = assign_groups(out["medicine_name"].tolist(), existing_groups)
    if "_base_given" in out.columns:
        given = out["_base_given"].map(_cell_text)
        out["base_drug_name"] = [g.upper() if g else e for g, e in zip(given, extracted)]
        out = out.drop(columns=["_base_given"])
    else:
        out["base_drug_name"] = extracted
    out = out[out["base_drug_name"] != "UNKNOWN"]

    # Same item twice (ignoring case) -> keep the first
    out = out[~out["medicine_name"].str.lower().duplicated()]

    ordered = ["medicine_name", "base_drug_name"]
    ordered += [f for f in NUMERIC_FIELDS if f in out.columns]
    ordered += [f for f in TRACKING_FIELDS if f in out.columns]
    ordered += [c for c in out.columns if c not in ordered]
    return out[ordered].reset_index(drop=True)
